"""把通用 Markdown 转成可发给客户的 Word。

覆盖报价、说明类文档用到的写法：一级/二级标题、段落、加粗、
有序和无序列表、表格。版式为 A4、微软雅黑、深蓝章节线。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

NAVY = "1B3A5F"
INK = "222222"
MUTED = "555555"
HEADER_BG = "1B3A5F"
ZEBRA = "F4F7FB"
LINE = "D0D7E2"
FONT = "微软雅黑"


def _rfonts(run) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)


def add_runs(paragraph, text: str, size: float, color: str, bold: bool = False) -> None:
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold or idx % 2 == 1
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        _rfonts(run)


def set_spacing(paragraph, before: int, after: int, line: int = 276, left: int = 0) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    if left:
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:left"), str(left))


def section_border(paragraph) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), NAVY)
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", "60"), ("bottom", "60"), ("left", "80"), ("right", "80")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LINE)
        borders.append(node)
    tbl_pr.append(borders)


def set_table_width(table, width_cm: float) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    twips = str(int(width_cm * 567))
    tbl_w.set(qn("w:w"), twips)
    tbl_w.set(qn("w:type"), "dxa")


def align_of(spec: str) -> WD_ALIGN_PARAGRAPH:
    spec = spec.strip()
    if spec.startswith(":") and spec.endswith(":"):
        return WD_ALIGN_PARAGRAPH.CENTER
    if spec.endswith(":"):
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], list[WD_ALIGN_PARAGRAPH], int]:
    raw_rows: list[list[str]] = []
    aligns: list[WD_ALIGN_PARAGRAPH] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if all(set(cell) <= set("-: ") and cell for cell in cells):
            aligns = [align_of(cell) for cell in cells]
        else:
            raw_rows.append(cells)
        idx += 1
    if not aligns:
        aligns = [WD_ALIGN_PARAGRAPH.LEFT] * len(raw_rows[0])
    width = max(len(row) for row in raw_rows)
    aligns = (aligns + [WD_ALIGN_PARAGRAPH.LEFT] * width)[:width]
    rows = [row + [""] * (width - len(row)) for row in raw_rows]
    return rows, aligns, idx


def add_table(doc: Document, rows: list[list[str]], aligns: list[WD_ALIGN_PARAGRAPH], width_cm: float) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    set_table_width(table, width_cm)
    set_table_borders(table)
    col_w = int(width_cm * 567 / len(rows[0]))
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Twips(col_w)
            set_cell_margins(cell)
            if r_idx == 0:
                shade_cell(cell, HEADER_BG)
            elif r_idx % 2 == 0:
                shade_cell(cell, ZEBRA)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = aligns[c_idx]
            set_spacing(paragraph, 20, 20, 240)
            color = "FFFFFF" if r_idx == 0 else INK
            add_runs(paragraph, text, 9, color, bold=(r_idx == 0))
    spacer = doc.add_paragraph()
    set_spacing(spacer, 0, 60, 20)


def build(md_text: str, out_path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    content_width = 17.0

    lines = [line.rstrip() for line in md_text.splitlines()]
    idx = 0
    seen_title = False

    def skip_blank() -> None:
        nonlocal idx
        while idx < len(lines) and not lines[idx].strip():
            idx += 1

    while idx < len(lines):
        skip_blank()
        if idx >= len(lines):
            break
        line = lines[idx].strip()
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            if not seen_title:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_spacing(paragraph, 0, 80, 276)
                add_runs(paragraph, line[2:].strip(), 18, NAVY, bold=True)
                seen_title = True
            else:
                set_spacing(paragraph, 220, 80, 276)
                section_border(paragraph)
                add_runs(paragraph, line[2:].strip(), 14, NAVY, bold=True)
            idx += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph()
            set_spacing(paragraph, 140, 40, 276)
            add_runs(paragraph, line[3:].strip(), 12, NAVY, bold=True)
            idx += 1
            continue
        if line.startswith("|"):
            rows, aligns, idx = parse_table(lines, idx)
            add_table(doc, rows, aligns, content_width)
            continue
        if re.match(r"^\d+\.\s+", line):
            items: list[str] = []
            while idx < len(lines) and re.match(r"^\d+\.\s+", lines[idx].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[idx].strip()))
                idx += 1
            for n, item in enumerate(items, start=1):
                paragraph = doc.add_paragraph()
                set_spacing(paragraph, 20, 20, 276, left=360)
                add_runs(paragraph, f"{n}. {item}", 11, INK)
            continue
        if line.startswith("- "):
            items = []
            while idx < len(lines) and lines[idx].strip().startswith("- "):
                items.append(lines[idx].strip()[2:])
                idx += 1
            for item in items:
                paragraph = doc.add_paragraph()
                set_spacing(paragraph, 20, 20, 276, left=360)
                add_runs(paragraph, f"· {item}", 11, INK)
            continue
        paragraph = doc.add_paragraph()
        centered = line.startswith("日期：") or line.startswith("费用：")
        if centered:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_spacing(paragraph, 0, 20, 276)
            add_runs(paragraph, line, 10.5, MUTED)
        else:
            set_spacing(paragraph, 40, 40, 300)
            add_runs(paragraph, line, 11, INK)
        idx += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    SRC = Path("/Users/shunyaoyin/Documents/ai-company/veo/doc/用户文档/停车结果识别算法报价方案.md")
    DST = SRC.with_suffix(".docx")
    build(SRC.read_text(encoding="utf-8"), DST)
    print(DST)
