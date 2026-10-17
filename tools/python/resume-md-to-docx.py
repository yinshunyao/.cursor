"""把求职简历 Markdown 转成与既有 Word 版式一致的 docx。

版式对齐「个人简历-202609.docx」：A4、微软雅黑、深蓝章节下划线、
时间 | 单位 | 角色 三分栏。只覆盖本简历用到的 Markdown（标题、表格、列表、加粗）。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1B3A5F"
BLUE = "2C5F8A"
INK = "222222"
MUTED = "555555"
FONT = "微软雅黑"


def _rfonts(run) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)


def add_run(paragraph, text: str, size: float, color: str, bold: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    _rfonts(run)


def set_spacing(paragraph, before: int, after: int, line: int, left: int = 0) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    for tag in ("w:spacing", "w:ind"):
        old = p_pr.find(qn(tag))
        if old is not None:
            p_pr.remove(old)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "0")
    if left:
        ind.set(qn("w:left"), str(left))
    p_pr.append(ind)


def section_border(paragraph) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), NAVY)
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def split_bold(text: str) -> list[tuple[str, bool]]:
    parts = re.split(r"\*\*(.+?)\*\*", text)
    return [(part, idx % 2 == 1) for idx, part in enumerate(parts) if part]


def is_triple(text: str) -> bool:
    return text.count("|") >= 2


def triple_parts(text: str) -> list[str]:
    return [part.strip() for part in text.split("|")]


def pad_label(label: str) -> str:
    if len(label) >= 4:
        return label
    gap = " " * (4 * (4 - len(label)))
    if len(label) == 2:
        return label[0] + gap + label[1]
    return label + gap


def build(md_text: str, out_path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)

    lines = [line.rstrip() for line in md_text.splitlines()]
    idx = 0
    in_entry = False

    def blank() -> None:
        nonlocal idx
        while idx < len(lines) and (not lines[idx].strip() or lines[idx].strip() == "---"):
            idx += 1

    blank()
    if not lines[idx].startswith("# "):
        raise ValueError("简历必须以一级标题开头")
    name = lines[idx][2:].strip()
    idx += 1
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(name_p, 0, 40, 440)
    add_run(name_p, "  ".join(name), 20, NAVY, bold=True)

    header_lines: list[str] = []
    blank()
    while idx < len(lines) and not lines[idx].startswith("## "):
        if lines[idx].strip() and lines[idx].strip() != "---":
            header_lines.append(lines[idx].strip())
        idx += 1

    for n, raw in enumerate(header_lines):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        is_tagline = n == len(header_lines) - 1 and "·" in raw and "|" not in raw
        if is_tagline:
            set_spacing(paragraph, 0, 40, 300)
            shown = re.sub(r"\s*·\s*", "  ·  ", raw)
            add_run(paragraph, shown, 10, BLUE, bold=True)
            continue
        set_spacing(paragraph, 0, 20, 300)
        color = INK if n == 0 else MUTED
        if "|" in raw:
            bits = [bit.strip() for bit in raw.split("|")]
            for bit_i, bit in enumerate(bits):
                if bit_i:
                    add_run(paragraph, "  |  ", 10.5, color)
                for text, bold in split_bold(bit):
                    add_run(paragraph, text, 10.5, color, bold=bold)
        else:
            for text, bold in split_bold(raw):
                add_run(paragraph, text, 10.5, color, bold=bold)

    while idx < len(lines):
        blank()
        if idx >= len(lines):
            break
        line = lines[idx].strip()
        if line.startswith("## "):
            in_entry = False
            paragraph = doc.add_paragraph()
            set_spacing(paragraph, 140, 60, 300)
            section_border(paragraph)
            add_run(paragraph, line[3:].strip(), 12.5, NAVY, bold=True)
            idx += 1
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
                if not all(set(cell) <= set("-: ") for cell in cells):
                    rows.append(cells)
                idx += 1
            for row in rows[1:]:
                label = pad_label(row[0])
                body = row[1] if len(row) > 1 else ""
                paragraph = doc.add_paragraph()
                set_spacing(paragraph, 0, 24, 290)
                add_run(paragraph, f"{label}：", 10.5, NAVY, bold=True)
                add_run(paragraph, body, 10.5, INK)
            continue
        if line.startswith("### "):
            in_entry = True
            title = line[4:].strip()
            paragraph = doc.add_paragraph()
            set_spacing(paragraph, 100, 30, 300)
            if is_triple(title):
                date, org, role = triple_parts(title)[:3]
                add_run(paragraph, date, 10.5, NAVY, bold=True)
                add_run(paragraph, "  |  ", 10.5, MUTED)
                add_run(paragraph, org, 10.5, INK, bold=True)
                add_run(paragraph, "  |  ", 10.5, MUTED)
                add_run(paragraph, role, 10.5, BLUE, bold=True)
            else:
                add_run(paragraph, title, 10.5, NAVY, bold=True)
            idx += 1
            continue
        if line.startswith("- "):
            body = line[2:].strip()
            paragraph = doc.add_paragraph()
            if is_triple(body):
                set_spacing(paragraph, 40, 20, 300)
                date, org, role = triple_parts(body)[:3]
                add_run(paragraph, date, 10.5, NAVY, bold=True)
                add_run(paragraph, "  |  ", 10.5, MUTED)
                add_run(paragraph, org, 10.5, INK, bold=True)
                add_run(paragraph, "  |  ", 10.5, MUTED)
                add_run(paragraph, role, 10.5, BLUE, bold=True)
            else:
                set_spacing(paragraph, 0, 24, 290, left=125 if in_entry else 0)
                add_run(paragraph, "•  ", 10.5, BLUE)
                for text, bold in split_bold(body):
                    add_run(paragraph, text, 10.5, INK, bold=bold)
            idx += 1
            continue
        if re.fullmatch(r"\*\*[^*]+\*\*", line):
            paragraph = doc.add_paragraph()
            set_spacing(paragraph, 40, 20, 290, left=125 if in_entry else 0)
            add_run(paragraph, line.strip("*") + "：", 10.5, NAVY, bold=True)
            idx += 1
            continue
        paragraph = doc.add_paragraph()
        if "·" in line and "|" not in line and not line.startswith("场景"):
            set_spacing(paragraph, 40, 0, 300)
            shown = re.sub(r"\s*·\s*", "    ·    ", line)
            add_run(paragraph, shown, 10.5, INK)
        else:
            set_spacing(paragraph, 0, 30, 290, left=125 if in_entry else 0)
            for text, bold in split_bold(line):
                add_run(paragraph, text, 10.5, INK, bold=bold)
        idx += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    SRC = Path(
        "/Users/shunyaoyin/Documents/ai-company/study/成都智臻善算科技公司/00 公司管理/法人相关信息/个人简历-项目经理-202610.md"
    )
    DST = SRC.with_suffix(".docx")
    build(SRC.read_text(encoding="utf-8"), DST)
    print(DST)
